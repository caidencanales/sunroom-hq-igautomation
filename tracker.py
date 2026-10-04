"""tracker.xlsx - one spreadsheet with everything.

Sheets:
  Dashboard  totals, reply rates per channel, today's numbers
  Leads      every lead: score, links, messages, what was sent when, replies
  Daily      sends + replies per day per channel
  Replies    just the people who answered

Type in the "Manual" column of the Leads sheet to override a lead:
  replied | won | skip | not interested     (synced back on the next run)
"""
import logging
import os
from datetime import datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import db
from config import TRACKER_PATH

log = logging.getLogger("sunroom")

# (header, db column, width)
LEAD_COLS = [
    ("Score", "score", 7), ("Status", "status", 11), ("Manual", "manual", 12), ("Company", "company", 24),
    ("Owner", "owner_name", 18), ("Website", "website", 26), ("Instagram", "instagram", 30),
    ("IG followers", "ig_followers", 11), ("Owner LinkedIn", "owner_linkedin", 32), ("Email", "email", 28),
    ("Why they fit", "score_reason", 45), ("Hook", "hook", 35), ("IG DM", "ig_message", 50),
    ("LinkedIn note", "li_note", 45), ("Email subject", "email_subject", 22), ("Email body", "email_body", 60),
    ("IG sent", "ig_sent_at", 17), ("LI sent", "li_sent_at", 17), ("Email sent", "email_sent_at", 17),
    ("Follow-up 1", "fu1_sent_at", 17), ("Follow-up 2", "fu2_sent_at", 17), ("Replied", "replied_at", 17),
    ("Reply via", "reply_channel", 10), ("Reply", "reply_snippet", 50), ("Source", "source", 28),
    ("Notes", "notes", 40), ("Added", "created_at", 17), ("Key", "key", 26),
]
MANUAL_MAP = {"replied": "replied", "won": "won", "skip": "done", "not interested": "done", "opted out": "opted_out"}
HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(bold=True, color="FFFFFF")
STATUS_FILL = {"replied": "C6F6D5", "won": "9AE6B4", "active": "FEFCBF", "opted_out": "FED7D7",
               "bounced": "FED7D7", "rejected": "EDF2F7"}


def sync_in():
    """Pull anything you typed in the Manual column back into the database."""
    if not os.path.exists(TRACKER_PATH):
        return
    try:
        ws = load_workbook(TRACKER_PATH, read_only=True)["Leads"]
    except Exception as e:
        log.warning(f"[tracker] couldn't read {TRACKER_PATH}: {e}")
        return
    headers = [cell.value for cell in next(ws.iter_rows(max_row=1))]
    i_manual, i_key = headers.index("Manual"), headers.index("Key")
    c = db.conn()
    changed = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        manual, key = row[i_manual], row[i_key]
        if not manual or not key:
            continue
        manual = str(manual).strip().lower()
        cur = c.execute("SELECT manual FROM leads WHERE key=?", (key,)).fetchone()
        if cur is None or cur["manual"] == manual:
            continue
        fields = {"manual": manual}
        if manual in MANUAL_MAP:
            fields["status"] = MANUAL_MAP[manual]
            if manual == "replied":
                fields.update(replied_at=db.now(), reply_channel="manual")
        db.update(c, key, **fields)
        changed += 1
    if changed:
        log.info(f"[tracker] synced {changed} manual updates")


def _header(ws, headers, widths=None):
    ws.append(headers)
    for i, cell in enumerate(ws[1], 1):
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = (widths or [18] * len(headers))[i - 1]
    ws.freeze_panes = "A2"


def export():
    c = db.conn()
    wb = Workbook()

    # Dashboard
    ws = wb.active
    ws.title = "Dashboard"
    total = lambda sql, *a: c.execute(sql, a).fetchone()[0]
    today = datetime.now().date().isoformat()
    rows = [
        ("Updated", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("", ""),
        ("Leads found", total("SELECT COUNT(*) FROM leads")),
        ("Qualified (score passed)", total("SELECT COUNT(*) FROM leads WHERE score IS NOT NULL AND status!='rejected'")),
        ("Rejected", total("SELECT COUNT(*) FROM leads WHERE status='rejected'")),
        ("Waiting to be contacted", db.queue_size(c)),
        ("", ""),
        ("Channel", "Sent (all time) / Sent today / Replies / Reply rate"),
    ]
    for label, col in (("Instagram", "ig_sent_at"), ("LinkedIn", "li_sent_at"), ("Email", "email_sent_at")):
        sent = total(f"SELECT COUNT(*) FROM leads WHERE {col} LIKE '20%'")
        sent_today = total(f"SELECT COUNT(*) FROM leads WHERE {col} LIKE ?", f"{today}%")
        replies = total(f"SELECT COUNT(*) FROM leads WHERE {col} LIKE '20%' AND status IN ('replied','won')")
        rows.append((label, f"{sent} / {sent_today} / {replies} / {replies / sent:.1%}" if sent else "0 / 0 / 0 / -"))
    rows += [
        ("", ""),
        ("Replied", total("SELECT COUNT(*) FROM leads WHERE status='replied'")),
        ("Won", total("SELECT COUNT(*) FROM leads WHERE status='won'")),
        ("Opted out", total("SELECT COUNT(*) FROM leads WHERE status='opted_out'")),
        ("Bounced", total("SELECT COUNT(*) FROM leads WHERE status='bounced'")),
    ]
    for r in rows:
        ws.append(r)
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 28, 50
    for cell in ws["A"]:
        cell.font = Font(bold=True)

    # Leads
    ws = wb.create_sheet("Leads")
    _header(ws, [h for h, _, _ in LEAD_COLS], [w for _, _, w in LEAD_COLS])
    cols = ", ".join(col for _, col, _ in LEAD_COLS)
    i_status = [col for _, col, _ in LEAD_COLS].index("status") + 1
    link_cols = {i + 1 for i, (_, col, _) in enumerate(LEAD_COLS) if col in ("website", "instagram", "owner_linkedin")}
    for r in c.execute(f"SELECT {cols} FROM leads ORDER BY (status='rejected'), score DESC"):
        ws.append(list(r))
        row = ws.max_row
        fill = STATUS_FILL.get(r["status"])
        if fill:
            ws.cell(row, i_status).fill = PatternFill("solid", fgColor=fill)
        for i in link_cols:
            cell = ws.cell(row, i)
            if cell.value and str(cell.value).startswith("http"):
                cell.hyperlink, cell.style = cell.value, "Hyperlink"
    ws.auto_filter.ref = ws.dimensions

    # Daily
    ws = wb.create_sheet("Daily")
    _header(ws, ["Date", "Instagram", "LinkedIn", "Email", "Follow-ups", "Replies", "New leads"])
    days = c.execute(
        "SELECT d FROM (SELECT substr(ig_sent_at,1,10) d FROM leads UNION SELECT substr(li_sent_at,1,10) FROM leads "
        "UNION SELECT substr(email_sent_at,1,10) FROM leads UNION SELECT substr(created_at,1,10) FROM leads) "
        "WHERE d LIKE '20%' ORDER BY d DESC"
    ).fetchall()
    for (d,) in days:
        n = lambda col: total(f"SELECT COUNT(*) FROM leads WHERE {col} LIKE ?", f"{d}%")
        ws.append([d, n("ig_sent_at"), n("li_sent_at"), n("email_sent_at"), n("fu1_sent_at") + n("fu2_sent_at"),
                   n("replied_at"), n("created_at")])

    # Replies
    ws = wb.create_sheet("Replies")
    _header(ws, ["Replied", "Via", "Status", "Company", "Owner", "Email", "Instagram", "Reply"],
            [17, 9, 11, 24, 18, 28, 30, 80])
    for r in c.execute("SELECT replied_at, reply_channel, status, company, owner_name, email, instagram, reply_snippet "
                       "FROM leads WHERE replied_at IS NOT NULL ORDER BY replied_at DESC"):
        ws.append(list(r))

    try:
        wb.save(TRACKER_PATH)
        path = TRACKER_PATH
    except PermissionError:  # the file is open in Excel
        path = TRACKER_PATH.replace(".xlsx", f"_{datetime.now():%Y%m%d_%H%M}.xlsx")
        wb.save(path)
        log.warning(f"[tracker] {TRACKER_PATH} is open in Excel - saved to {path} instead")
    log.info(f"[tracker] saved {path}")
    return path
